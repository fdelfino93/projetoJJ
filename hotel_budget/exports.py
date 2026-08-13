"""Exportação de relatórios em CSV, PDF e PowerPoint.

Os gráficos embutidos aqui vêm do `graficos` em versão Matplotlib (PNG), já
que Plotly serve à tela e não a arquivos estáticos.
"""

import io
from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import analise
import cadastro
import graficos
from config import (
    COR_CRITICO,
    COR_DESPESA,
    COR_OK,
    COR_RECEITA,
    COR_TEXTO,
    HOTEL_NOME,
    SISTEMA_NOME,
    VERSAO,
    mes_nome,
)
from utils import moeda, validar_mes, validar_tipo

_AZUL = colors.HexColor(COR_RECEITA)
_LARANJA = colors.HexColor(COR_DESPESA)
_VERDE = colors.HexColor(COR_OK)
_VERMELHO = colors.HexColor(COR_CRITICO)
_CINZA_CLARO = colors.HexColor("#f4f4f2")
_TINTA = colors.HexColor(COR_TEXTO)

_RGB_AZUL = RGBColor(0x2A, 0x78, 0xD6)
_RGB_TINTA = RGBColor(0x0B, 0x0B, 0x0B)
_RGB_SUAVE = RGBColor(0x52, 0x51, 0x4E)


def _rotulo(tipo: str) -> str:
    return "Receitas" if tipo == "receita" else "Despesas"


# --------------------------------------------------------------------- CSV


def gerar_csv_movimentacoes(tipo: str, mes: Optional[str] = None) -> str:
    """Lançamentos tratados, em CSV com ponto e vírgula (padrão Excel BR)."""
    tipo = validar_tipo(tipo)
    df = cadastro.listar_movimentacoes(tipo, mes)
    linhas = ["id;data;categoria;descricao;valor"]
    for linha in df.itertuples():
        data = linha.data.strftime("%Y-%m-%d")
        descricao = str(linha.descricao).replace(";", ",")
        linhas.append(f"{linha.id};{data};{linha.categoria};{descricao};{linha.valor:.2f}")
    linhas.append(f";;;TOTAL;{df['valor'].sum():.2f}")
    return "\n".join(linhas)


def gerar_csv_relatorio(mes: str) -> str:
    """Relatório mensal completo em CSV, pronto para abrir no Excel."""
    mes = validar_mes(mes)
    relatorio = analise.relatorio_mensal(mes)
    resumo = relatorio["resumo"]

    linhas = [
        "RELATORIO MENSAL",
        f"Hotel;{HOTEL_NOME}",
        f"Mes;{mes_nome(mes)}",
        "",
        "RESUMO;VALOR",
        f"Receitas planejadas;{resumo['receita_planejada']:.2f}",
        f"Receitas realizadas;{resumo['receita_realizada']:.2f}",
        f"Despesas planejadas;{resumo['despesa_planejada']:.2f}",
        f"Despesas realizadas;{resumo['despesa_realizada']:.2f}",
        f"Saldo planejado;{resumo['saldo_planejado']:.2f}",
        f"Saldo realizado;{resumo['saldo_realizado']:.2f}",
        f"Execucao das receitas (%);{resumo['execucao_receita']:.2f}",
        f"Execucao das despesas (%);{resumo['execucao_despesa']:.2f}",
        f"Margem (%);{resumo['margem']:.2f}",
        "",
        "TIPO;CATEGORIA;PLANEJADO;REALIZADO;VARIACAO;EXECUCAO (%);PARTICIPACAO (%);SITUACAO",
    ]
    for tipo in ("receita", "despesa"):
        for linha in relatorio[f"{tipo}s"].itertuples():
            linhas.append(
                f"{tipo};{linha.categoria};{linha.planejado:.2f};{linha.realizado:.2f};"
                f"{linha.variacao:.2f};{linha.execucao:.2f};{linha.participacao:.2f};{linha.situacao}"
            )

    if relatorio["alertas"]:
        linhas += ["", "ALERTAS;SITUACAO"]
        linhas += [f"{a['mensagem']};{a['situacao']}" for a in relatorio["alertas"]]
    return "\n".join(linhas)


# --------------------------------------------------------------------- PDF


def _estilos():
    """Conjunto de estilos de texto usado nos documentos PDF."""
    base = getSampleStyleSheet()
    return {
        "base": base,
        "titulo": ParagraphStyle("Titulo", parent=base["Title"], fontSize=18, textColor=_TINTA, alignment=0),
        "subtitulo": ParagraphStyle("Subtitulo", parent=base["Normal"], fontSize=10, textColor=colors.grey),
        "secao": ParagraphStyle("Secao", parent=base["Heading2"], fontSize=13, textColor=_TINTA, spaceAfter=4),
        "cabecalho": ParagraphStyle("Cab", parent=base["Normal"], fontSize=9, textColor=colors.white),
        "celula": ParagraphStyle("Celula", parent=base["Normal"], fontSize=9),
        "celula_dir": ParagraphStyle("CelulaDir", parent=base["Normal"], fontSize=9, alignment=TA_RIGHT),
        "centro": ParagraphStyle("Centro", parent=base["Normal"], fontSize=11, alignment=TA_CENTER),
    }


def _tabela(dados: list[list], cor_cabecalho=_AZUL, larguras=None) -> Table:
    """Monta uma tabela com cabeçalho colorido e zebra nas linhas."""
    estilos = _estilos()
    linhas = [[Paragraph(str(c), estilos["cabecalho"]) for c in dados[0]]]
    for linha in dados[1:]:
        celulas = [Paragraph(str(linha[0]), estilos["celula"])]
        celulas += [Paragraph(str(c), estilos["celula_dir"]) for c in linha[1:]]
        linhas.append(celulas)

    tabela = Table(linhas, colWidths=larguras, repeatRows=1)
    tabela.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), cor_cabecalho),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, _CINZA_CLARO]),
                ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#e1e0d9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return tabela


def _imagem(png: Optional[bytes], largura_mm: float = 165) -> Optional[Image]:
    """Converte um PNG em elemento de imagem proporcional ao PDF."""
    if not png:
        return None
    imagem = Image(io.BytesIO(png))
    proporcao = imagem.imageHeight / imagem.imageWidth
    imagem.drawWidth = largura_mm * mm
    imagem.drawHeight = largura_mm * mm * proporcao
    return imagem


def _cabecalho_pdf(estilos, titulo: str, subtitulo: str) -> list:
    return [
        Paragraph(titulo, estilos["titulo"]),
        Paragraph(subtitulo, estilos["subtitulo"]),
        Spacer(1, 6 * mm),
    ]


def gerar_pdf_movimentacoes(tipo: str, mes: Optional[str] = None) -> bytes:
    """Extrato de lançamentos em PDF."""
    tipo = validar_tipo(tipo)
    estilos = _estilos()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, title=f"{SISTEMA_NOME} — {_rotulo(tipo)}",
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
    )

    df = cadastro.listar_movimentacoes(tipo, mes)
    periodo = mes_nome(mes) if mes else "Período completo"
    elementos = _cabecalho_pdf(estilos, _rotulo(tipo), f"{HOTEL_NOME} — {periodo}")

    if df.empty:
        elementos.append(Paragraph("Nenhum lançamento encontrado.", estilos["celula"]))
    else:
        dados = [["Data", "Categoria", "Descrição", "Valor"]]
        for linha in df.itertuples():
            dados.append([
                linha.data.strftime("%d/%m/%Y"),
                linha.categoria,
                str(linha.descricao) or "—",
                moeda(linha.valor),
            ])
        dados.append(["", "", "TOTAL", moeda(df["valor"].sum())])
        cor = _AZUL if tipo == "receita" else _LARANJA
        elementos.append(_tabela(dados, cor, larguras=[24 * mm, 48 * mm, 68 * mm, 34 * mm]))

    doc.build(elementos)
    return buf.getvalue()


def gerar_pdf_relatorio(mes: str) -> bytes:
    """Relatório mensal em PDF, com resumo, alertas, gráficos e comparativos."""
    mes = validar_mes(mes)
    estilos = _estilos()
    relatorio = analise.relatorio_mensal(mes)
    resumo = relatorio["resumo"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, title=f"{SISTEMA_NOME} — Relatório de {mes_nome(mes)}",
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
    )

    elementos = _cabecalho_pdf(
        estilos,
        f"Relatório Mensal — {mes_nome(mes)}",
        f"{HOTEL_NOME} · {SISTEMA_NOME} v{VERSAO}",
    )

    # Resumo em três colunas.
    saldo_cor = COR_OK if resumo["saldo_realizado"] >= 0 else COR_CRITICO
    resumo_tabela = Table(
        [
            [
                Paragraph("<b>Receitas</b>", estilos["centro"]),
                Paragraph("<b>Despesas</b>", estilos["centro"]),
                Paragraph("<b>Saldo</b>", estilos["centro"]),
            ],
            [
                Paragraph(f"<font color='{COR_RECEITA}'>{moeda(resumo['receita_realizada'])}</font>", estilos["centro"]),
                Paragraph(f"<font color='{COR_DESPESA}'>{moeda(resumo['despesa_realizada'])}</font>", estilos["centro"]),
                Paragraph(f"<font color='{saldo_cor}'>{moeda(resumo['saldo_realizado'])}</font>", estilos["centro"]),
            ],
            [
                Paragraph(f"{resumo['execucao_receita']:.1f}% do planejado", estilos["centro"]),
                Paragraph(f"{resumo['execucao_despesa']:.1f}% do planejado", estilos["centro"]),
                Paragraph(f"margem de {resumo['margem']:.1f}%", estilos["centro"]),
            ],
        ],
        colWidths=[58 * mm] * 3,
    )
    resumo_tabela.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _CINZA_CLARO),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#e1e0d9")),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e1e0d9")),
        ])
    )
    elementos += [resumo_tabela, Spacer(1, 7 * mm)]

    if relatorio["alertas"]:
        elementos.append(Paragraph("Pontos de atenção", estilos["secao"]))
        for alerta in relatorio["alertas"][:6]:
            cor = COR_CRITICO if alerta["situacao"] == "critico" else "#b07a00"
            elementos.append(
                Paragraph(f"<font color='{cor}'>•</font> {alerta['mensagem']}", estilos["celula"])
            )
        elementos.append(Spacer(1, 6 * mm))

    for tipo in ("receita", "despesa"):
        elementos.append(Paragraph(f"{_rotulo(tipo)} — Planejado x Realizado", estilos["secao"]))
        dados = [["Categoria", "Planejado", "Realizado", "Variação", "Execução"]]
        for linha in relatorio[f"{tipo}s"].itertuples():
            dados.append([
                linha.categoria,
                moeda(linha.planejado),
                moeda(linha.realizado),
                moeda(linha.variacao),
                f"{linha.execucao:.1f}%",
            ])
        cor = _AZUL if tipo == "receita" else _LARANJA
        elementos.append(_tabela(dados, cor, larguras=[52 * mm, 30 * mm, 30 * mm, 30 * mm, 22 * mm]))
        elementos.append(Spacer(1, 4 * mm))
        imagem = _imagem(graficos.png_planejado_realizado(tipo, mes))
        if imagem:
            elementos += [imagem, Spacer(1, 7 * mm)]

    doc.build(elementos)
    return buf.getvalue()


# ---------------------------------------------------------------- PowerPoint


def _apresentacao() -> Presentation:
    """Apresentação widescreen 16:9 em branco."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _texto(slide, x, y, w, h, linhas, cor=_RGB_TINTA, negrito=False, centralizado=False):
    """Insere uma caixa de texto com uma linha por tupla (texto, tamanho)."""
    caixa = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    quadro = caixa.text_frame
    quadro.word_wrap = True
    for i, (conteudo, tamanho) in enumerate(linhas):
        paragrafo = quadro.paragraphs[0] if i == 0 else quadro.add_paragraph()
        paragrafo.text = conteudo
        paragrafo.font.size = Pt(tamanho)
        paragrafo.font.bold = negrito
        paragrafo.font.color.rgb = cor
        if centralizado:
            paragrafo.alignment = PP_ALIGN.CENTER
    return caixa


def _slide_capa(prs, titulo: str, subtitulo: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _texto(slide, 1.0, 2.4, 11.3, 1.2, [(SISTEMA_NOME, 40)], _RGB_AZUL, True, True)
    _texto(slide, 1.0, 3.5, 11.3, 1.0, [(titulo, 26)], _RGB_TINTA, True, True)
    _texto(slide, 1.0, 4.4, 11.3, 1.4, [(subtitulo, 16), (f"versão {VERSAO}", 12)], _RGB_SUAVE, False, True)
    return slide


def _slide_titulo(prs, titulo: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _texto(slide, 0.7, 0.4, 11.9, 0.9, [(titulo, 28)], _RGB_AZUL, True)
    return slide


def _slide_imagem(prs, titulo: str, png: Optional[bytes]):
    """Slide com um gráfico centralizado."""
    slide = _slide_titulo(prs, titulo)
    if png:
        slide.shapes.add_picture(io.BytesIO(png), Inches(1.6), Inches(1.5), width=Inches(10.1))
    else:
        _texto(slide, 0.7, 3.0, 11.9, 1.0, [("Sem dados para o período.", 18)], _RGB_SUAVE)
    return slide


def _slide_indicadores(prs, titulo: str, itens: list[tuple[str, str]]):
    """Slide com pares rótulo/valor em duas colunas."""
    slide = _slide_titulo(prs, titulo)
    metade = (len(itens) + 1) // 2
    for coluna, grupo in enumerate((itens[:metade], itens[metade:])):
        x = 0.9 + coluna * 6.2
        for linha, (rotulo, valor) in enumerate(grupo):
            y = 1.6 + linha * 1.15
            _texto(slide, x, y, 5.6, 0.4, [(rotulo, 13)], _RGB_SUAVE)
            _texto(slide, x, y + 0.34, 5.6, 0.6, [(valor, 24)], _RGB_TINTA, True)
    return slide


def gerar_pptx_relatorio(mes: str) -> bytes:
    """Apresentação do fechamento mensal."""
    mes = validar_mes(mes)
    resumo = analise.resumo_mes(mes)
    projecao = analise.projecao_fechamento(mes)
    alertas = analise.alertas(mes)

    prs = _apresentacao()
    _slide_capa(prs, f"Fechamento de {mes_nome(mes)}", HOTEL_NOME)
    _slide_indicadores(
        prs,
        f"Resumo de {mes_nome(mes)}",
        [
            ("Receitas realizadas", moeda(resumo["receita_realizada"])),
            ("Despesas realizadas", moeda(resumo["despesa_realizada"])),
            ("Saldo do mês", moeda(resumo["saldo_realizado"])),
            ("Execução das receitas", f"{resumo['execucao_receita']:.1f}%"),
            ("Execução das despesas", f"{resumo['execucao_despesa']:.1f}%"),
            ("Margem", f"{resumo['margem']:.1f}%"),
            ("Lançamentos no mês", f"{resumo['qtd_receitas'] + resumo['qtd_despesas']}"),
            ("Saldo projetado", moeda(projecao["saldo_projetado"])),
        ],
    )
    _slide_imagem(prs, "Resultado do mês", graficos.png_resultado_mes(mes))
    _slide_imagem(prs, "Receitas — Planejado x Realizado", graficos.png_planejado_realizado("receita", mes))
    _slide_imagem(prs, "Despesas — Planejado x Realizado", graficos.png_planejado_realizado("despesa", mes))
    _slide_imagem(prs, "Despesas por categoria", graficos.png_distribuicao("despesa", mes))

    if alertas:
        slide = _slide_titulo(prs, "Pontos de atenção")
        _texto(
            slide, 0.9, 1.6, 11.5, 5.0,
            [(f"• {a['mensagem']}", 16) for a in alertas[:8]],
            _RGB_TINTA,
        )

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def gerar_pptx_executivo(mes: str) -> bytes:
    """Apresentação curta para a diretoria: tendência e números-chave."""
    mes = validar_mes(mes)
    resumo = analise.resumo_mes(mes)

    prs = _apresentacao()
    _slide_capa(prs, f"Relatório Executivo — {mes_nome(mes)}", HOTEL_NOME)
    _slide_imagem(prs, "Evolução mensal", graficos.png_evolucao_mensal())
    _slide_indicadores(
        prs,
        "Números do mês",
        [
            ("Receitas realizadas", moeda(resumo["receita_realizada"])),
            ("Despesas realizadas", moeda(resumo["despesa_realizada"])),
            ("Saldo do mês", moeda(resumo["saldo_realizado"])),
            ("Margem", f"{resumo['margem']:.1f}%"),
        ],
    )
    _slide_imagem(prs, "Receitas por categoria", graficos.png_distribuicao("receita", mes))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _mes_anterior(mes: str) -> str:
    """Mês anterior no formato AAAA-MM."""
    ano, num = int(mes[:4]), int(mes[5:7])
    return f"{ano - 1}-12" if num == 1 else f"{ano}-{num - 1:02d}"


def gerar_pptx_comparativo(mes: str) -> bytes:
    """Apresentação comparando o mês escolhido com o mês anterior."""
    mes = validar_mes(mes)
    anterior = _mes_anterior(mes)
    atual = analise.resumo_mes(mes)
    passado = analise.resumo_mes(anterior)

    def variacao(agora: float, antes: float) -> str:
        if not antes:
            return "—"
        return f"{((agora - antes) / abs(antes)) * 100:+.1f}%"

    prs = _apresentacao()
    _slide_capa(prs, f"{mes_nome(mes)} vs {mes_nome(anterior)}", HOTEL_NOME)
    _slide_indicadores(
        prs,
        "Comparativo de resultados",
        [
            ("Receitas realizadas", f"{moeda(atual['receita_realizada'])}  ({variacao(atual['receita_realizada'], passado['receita_realizada'])})"),
            ("Despesas realizadas", f"{moeda(atual['despesa_realizada'])}  ({variacao(atual['despesa_realizada'], passado['despesa_realizada'])})"),
            ("Saldo", f"{moeda(atual['saldo_realizado'])}  ({variacao(atual['saldo_realizado'], passado['saldo_realizado'])})"),
            ("Margem", f"{atual['margem']:.1f}%  (antes {passado['margem']:.1f}%)"),
        ],
    )
    _slide_imagem(prs, "Evolução mensal", graficos.png_evolucao_mensal())
    _slide_imagem(prs, f"Despesas — Planejado x Realizado ({mes_nome(mes)})", graficos.png_planejado_realizado("despesa", mes))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
